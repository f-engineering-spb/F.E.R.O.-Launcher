"""Воспроизводимый базовый сценарий Excel: синтетическая книга, без данных клиента.

Проверяет инварианты маршрута (не конкретные значения замеров):
- список листов без загрузки книги;
- первый запрошенный лист готовит всю книгу;
- повтор — чистый дисковый кэш-хит без openpyxl;
- лимит строк применяется с флагом limited в метаданных.
Запуск: C:\\Python314\\python.exe -m unittest tests.test_excel_baseline -v
"""
from __future__ import annotations

import sys
import tempfile
import time
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl

from app.rendering import engine_excel


def _synthetic_book(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "First"
    for r in range(1, 51):
        for c in range(1, 6):
            ws.cell(row=r, column=c, value=r * 10 + c)
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = 1
    ws2.merge_cells("A2:B3")
    big = wb.create_sheet("Big")
    for r in range(1, engine_excel.MAX_ROWS + 101):
        big.cell(row=r, column=1, value=r)
    wb.save(str(path))


class ExcelBaselineScenario(unittest.TestCase):
    def test_baseline_invariants(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src = tmp / "syn.xlsx"
            t0 = time.perf_counter()
            _synthetic_book(src)
            build_ms = (time.perf_counter() - t0) * 1000
            cache = tmp / "cache"
            t1 = time.perf_counter()
            names = engine_excel.get_sheet_names(src)
            names_ms = (time.perf_counter() - t1) * 1000
            self.assertEqual(len(names), 3)
            t2 = time.perf_counter()
            metas = engine_excel.ensure_workbook_rendered(src, cache)
            first_ms = (time.perf_counter() - t2) * 1000
            self.assertEqual(len(metas), 3)
            self.assertEqual(len(list(cache.glob("sheet-*-*.html"))), 3)
            big = next(m for m in metas if m["rows"] > engine_excel.MAX_ROWS - 100)
            self.assertTrue(big["limited"])
            t3 = time.perf_counter()
            metas2 = engine_excel.ensure_workbook_rendered(src, cache)
            re_ms = (time.perf_counter() - t3) * 1000
            self.assertEqual(len(metas2), 3)
            print(f"\n[excel-baseline] build={build_ms:.0f}ms names={names_ms:.1f}ms "
                  f"first={first_ms:.0f}ms repeat={re_ms:.1f}ms")


class ExcelBaselineLosses(unittest.TestCase):
    """E03: синтетические проверки текущих потерь (без исправлений).

    Каждый тест описывает наблюдаемое поведение движка на искусственной
    книге без данных клиента:
    - формула без сохранённого значения показывает пустую клетку;
    - процентный/денежный форматы не применяются (видно сырое число);
    - превышение лимита строк помечается limited, но предупреждение
      в HTML отсутствует;
    - превышение лимита столбцов обрезается молча (limited=False);
    - намеренно скрытый столбец исключается без предупреждения
      (само скрытие ошибкой не считается).
    """

    def _render_one(self, tmp: Path, name: str, build) -> tuple[dict, str]:
        src = tmp / name
        build(src)
        cache = tmp / (name + ".cache")
        metas = engine_excel.ensure_workbook_rendered(src, cache)
        self.assertEqual(len(metas), 1)
        page = (cache / "sheet-1-v10.html").read_text(encoding="utf-8")
        return metas[0], page

    def test_formula_without_cached_value_is_empty(self):
        def build(src: Path) -> None:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "S"
            ws["A1"] = 7
            ws["A2"] = "=A1*2"
            wb.save(str(src))

        with tempfile.TemporaryDirectory() as tmp:
            _, page = self._render_one(Path(tmp), "formula.xlsx", build)
            self.assertNotIn(">14<", page)
            self.assertIn("<td", page)

    def test_percent_and_money_formats_not_applied(self):
        def build(src: Path) -> None:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "S"
            ws["B1"] = 0.5
            ws["B1"].number_format = "0%"
            ws["C1"] = 1234.5
            ws["C1"].number_format = "#,##0.00"
            wb.save(str(src))

        with tempfile.TemporaryDirectory() as tmp:
            _, page = self._render_one(Path(tmp), "formats.xlsx", build)
            self.assertIn(">0.5<", page)
            self.assertNotIn(">50%<", page)
            self.assertIn(">1234.5<", page)
            self.assertNotIn(">1,234.50<", page)

    def test_row_limit_flag_without_visible_warning(self):
        def build(src: Path) -> None:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Big"
            for r in range(1, engine_excel.MAX_ROWS + 11):
                ws.cell(row=r, column=1, value=r)
            wb.save(str(src))

        with tempfile.TemporaryDirectory() as tmp:
            meta, page = self._render_one(Path(tmp), "rows.xlsx", build)
            self.assertTrue(meta["limited"])
            self.assertNotIn('class="notice"', page)

    def test_column_limit_is_silent(self):
        def build(src: Path) -> None:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Wide"
            for c in range(1, engine_excel.MAX_COLS + 11):
                ws.cell(row=1, column=c, value=c)
            wb.save(str(src))

        with tempfile.TemporaryDirectory() as tmp:
            meta, page = self._render_one(Path(tmp), "wide.xlsx", build)
            self.assertEqual(meta["columns"], engine_excel.MAX_COLS)
            self.assertFalse(meta["limited"])
            self.assertNotIn(">101<", page)
            self.assertNotIn('class="notice"', page)

    def test_hidden_column_excluded_without_warning(self):
        def build(src: Path) -> None:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "H"
            ws["A1"] = "a"
            ws["B1"] = "b-hidden"
            ws["C1"] = "c"
            ws.column_dimensions["B"].hidden = True
            wb.save(str(src))

        with tempfile.TemporaryDirectory() as tmp:
            meta, page = self._render_one(Path(tmp), "hidden.xlsx", build)
            self.assertEqual(meta["columns"], 2)
            self.assertNotIn("b-hidden", page)
            self.assertNotIn('class="notice"', page)


if __name__ == "__main__":
    unittest.main()
