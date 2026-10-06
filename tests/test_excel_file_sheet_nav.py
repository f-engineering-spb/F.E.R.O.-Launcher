"""Tests for EXCEL-RESTORE-FILE-SHEET-NAV-001 contract:
1. One workbook = one card in ribbon with sheet buttons.
2. Sheets switch via buttons inside the card and preserve selection.
3. Fast sheet inspection and fast H-fast HTML per sheet.
4. Context menu target points to original Excel workbook everywhere.
5. Thumbnail zoom scaling allows 2x enlargement.
"""

import json
import re
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.backend import excel_speed  # noqa: E402
from app.backend import server  # noqa: E402


class TestExcelFileSheetNav(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.tmp_dir = Path(self.tmp.name)
        self.book_path = self.tmp_dir / "test_book.xlsx"
        self._create_multisheet_book(self.book_path)

    def tearDown(self):
        self.tmp.cleanup()

    def _create_multisheet_book(self, path: Path):
        wb = openpyxl.Workbook()
        ws0 = wb.active
        ws0.title = "Смета"
        ws0.append(["Код", "Наименование", "Сумма"])
        for r in range(1, 21):
            ws0.append([f"K-{r}", f"Позиция {r}", r * 1000])

        ws1 = wb.create_sheet(title="Материалы")
        ws1.append(["Артикул", "Материал", "Кол-во"])
        for r in range(1, 15):
            ws1.append([f"M-{r}", f"Бетон М{r}", r * 10])

        ws2 = wb.create_sheet(title="Итого")
        ws2.append(["Показатель", "Значение"])
        ws2.append(["Всего", 123456])

        wb.save(path)
        wb.close()

    def test_workbook_fast_metadata_has_all_sheets(self):
        """Backend returns sheet list and fast sheet 0 thumbnail URL."""
        cache_dir = self.tmp_dir / "cache"
        wb_info = excel_speed.get_fast_workbook(self.book_path, cache_dir)
        self.assertEqual(wb_info["name"], "test_book.xlsx")
        self.assertEqual(len(wb_info["sheets"]), 3)
        self.assertEqual(wb_info["sheets"][0]["name"], "Смета")
        self.assertEqual(wb_info["sheets"][1]["name"], "Материалы")
        self.assertEqual(wb_info["sheets"][2]["name"], "Итого")
        self.assertIn("/cache/", wb_info["thumbnailUrl"])
        self.assertTrue(wb_info["thumbnailUrl"].endswith("preview.html"))
        self.assertTrue(Path(cache_dir).exists())

    def test_per_sheet_html_caching_and_switch(self):
        """Each sheet gets its own deterministic cache key and fast rendering."""
        cache_dir = self.tmp_dir / "cache"
        sheet0 = excel_speed.get_fast_html(self.book_path, cache_dir, sheet_index=0)
        sheet1 = excel_speed.get_fast_html(self.book_path, cache_dir, sheet_index=1)
        self.assertNotEqual(sheet0["path"], sheet1["path"])
        self.assertEqual(sheet0["sheet"], "Смета")
        self.assertEqual(sheet1["sheet"], "Материалы")
        content0 = Path(sheet0["path"]).read_text(encoding="utf-8")
        content1 = Path(sheet1["path"]).read_text(encoding="utf-8")
        self.assertIn("Позиция 1", content0)
        self.assertIn("Бетон М1", content1)

        # Warm hit for sheet1
        sheet1_warm = excel_speed.get_fast_html(self.book_path, cache_dir, sheet_index=1)
        self.assertTrue(sheet1_warm["cacheHit"])
        self.assertEqual(sheet1_warm["engineMs"], 0.0)

    def test_server_excel_workbook_and_sheet_preview_handlers(self):
        """server.excel_workbook_preview and server.excel_sheet_preview contract."""
        info = server.excel_workbook_preview(self.book_path)
        self.assertEqual(info["name"], "test_book.xlsx")
        self.assertEqual(len(info["sheets"]), 3)
        self.assertEqual(info["sourceType"], "XLSX")

        sheet_res = server.excel_sheet_preview(self.book_path, sheet_index=1)
        self.assertEqual(sheet_res["index"], 1)
        self.assertEqual(sheet_res["name"], "Материалы")
        self.assertIn("/cache/excel/", sheet_res["url"])

    def test_frontend_card_structure_static_contract(self):
        """Verify frontend app.js creates sheet buttons inside card, syncs sheet index, and zooms to 10."""
        app_js = (ROOT / "app" / "frontend" / "app.js").read_text(encoding="utf-8")

        # 1. Excel card has sheet buttons container
        self.assertIn('sheetsBar.className = "excel-card-sheets-bar"', app_js)
        self.assertIn('sBtn.className = "excel-card-sheet-btn"', app_js)

        # 2. Sheet click updates activeSheetIndex and switches sheet in card iframe
        self.assertIn("workbook.activeSheetIndex = sheet.index", app_js)
        self.assertIn("frame.src = sheet.url", app_js)

        # 3. activateExcelWorkbook accepts targetSheetIndex
        self.assertIn("async function activateExcelWorkbook(index, targetSheetIndex = 0)", app_js)

        # 4. Sheet switch in viewer syncs back to ribbon card
        self.assertIn("cardThumb.querySelectorAll(\".excel-card-sheet-btn\")", app_js)

        # 5. zoomThumbs supports up to scale 10 (2x old limit 5)
        self.assertIn("Math.min(10,", app_js)

    def test_frontend_styles_card_and_zoom_contract(self):
        """Verify CSS contains styles for sheets bar and doubled max widths."""
        css_text = (ROOT / "app" / "frontend" / "styles.css").read_text(encoding="utf-8")

        self.assertIn(".excel-card-sheets-bar", css_text)
        self.assertIn(".excel-card-sheet-btn", css_text)
        self.assertIn(".excel-card-sheet-btn.active", css_text)

        # Doubled limits: 960px max-height for thumb image, 1720px for thumb-wrap
        self.assertIn("calc(960px * var(--thumb-scale, 1))", css_text)
        self.assertIn("calc(1720px * var(--thumb-scale, 1))", css_text)


if __name__ == "__main__":
    unittest.main()
