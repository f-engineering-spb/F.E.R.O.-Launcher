"""Регрессия DOCX→HTML: порядок, нумерация, ссылки, колонтитулы, кэш.

Только искусственные DOCX, собранные кодом теста. Никакого содержимого договора.
Запуск: C:\\Python314\\python.exe -m unittest tests.test_word_preview -v
(из корня проекта, PYTHONPATH=корень проекта)
"""
from __future__ import annotations

import base64
import concurrent.futures
import io
import json
import re
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from docx.styles.style import WD_STYLE_TYPE

from app.rendering import engine_word
from app.rendering.engine_word import (
    DOCX_HTML_RENDER_VERSION,
    docx_to_html_string,
    ensure_docx_preview,
)

# Синтетический PNG 1x1 (генерируется здесь, не из договора).
PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def _png_variant(marker: bytes) -> bytes:
    """Другой валидный PNG: добавляется tEXt-чанк с меткой (для тестов подмены)."""
    import binascii
    import struct
    body, iend = PIXEL_PNG[:-12], PIXEL_PNG[-12:]
    assert iend[4:8] == b"IEND"
    data = b"Comment\x00" + marker
    chunk = struct.pack(">I", len(data)) + b"tEXt" + data
    chunk += struct.pack(">I", binascii.crc32(b"tEXt" + data) & 0xFFFFFFFF)
    return body + chunk + iend


def _set_num_pr(paragraph, num_id: int, ilvl: int = 0) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    old = pPr.find(qn("w:numPr"))
    if old is not None:
        pPr.remove(old)
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), str(ilvl))
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numPr.append(il)
    numPr.append(nid)
    pPr.append(numPr)


def _add_numbering(doc, abstract_id: int, levels, num_id: int, overrides=None,
                   restarts=None, override_levels=None) -> None:
    """levels: [(start, numFmt, lvlText), ...].

    restarts: {ilvl: lvlRestart-val}. overrides: {ilvl: startOverride}.
    override_levels: {ilvl: (numFmt, lvlText)} — вложенный w:lvl в lvlOverride.
    """
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    lvl_xml = ""
    for i, (s, f, t) in enumerate(levels):
        restart = ""
        if restarts and i in restarts:
            restart = f'<w:lvlRestart w:val="{restarts[i]}"/>'
        lvl_xml += (f'<w:lvl w:ilvl="{i}"><w:start w:val="{s}"/>'
                    f'<w:numFmt w:val="{f}"/><w:lvlText w:val="{t}"/>{restart}</w:lvl>')
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        f'<w:multiLevelType w:val="multilevel"/>{lvl_xml}</w:abstractNum>')
    numbering.append(abstract)
    ov = ""
    for k, v in (overrides or {}).items():
        nested = ""
        if override_levels and k in override_levels:
            fmt, text = override_levels[k]
            nested = (f'<w:lvl w:ilvl="{k}"><w:start w:val="1"/>'
                      f'<w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/></w:lvl>')
        ov += (f'<w:lvlOverride w:ilvl="{k}"><w:startOverride w:val="{v}"/>' +
               nested + "</w:lvlOverride>")
    num = parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}">'
        f'<w:abstractNumId w:val="{abstract_id}"/>{ov}</w:num>')
    numbering.append(num)


def _set_num_id_only(paragraph, num_id: int) -> None:
    """Частичный прямой numPr: только numId, без ilvl (ilvl доберётся из стиля)."""
    pPr = paragraph._p.get_or_add_pPr()
    old = pPr.find(qn("w:numPr"))
    if old is not None:
        pPr.remove(old)
    numPr = OxmlElement("w:numPr")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numPr.append(nid)
    pPr.append(numPr)


def _add_break(run) -> None:
    run._r.append(OxmlElement("w:br"))


def _set_run_off(run, tag: str, val: str) -> None:
    rPr = run._r.get_or_add_rPr()
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), val)
    rPr.append(el)


def _add_hyperlink(paragraph, url: str, text: str):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    txt = OxmlElement("w:t")
    txt.text = text
    run.append(txt)
    link.append(run)
    paragraph._p.append(link)


def _add_anchor_link(paragraph, anchor: str, text: str):
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    txt = OxmlElement("w:t")
    txt.text = text
    run.append(txt)
    link.append(run)
    paragraph._p.append(link)
    # Якорь-назначение: закладка в другом абзаце добавит id="bm-...".


def _add_bookmark(paragraph, name: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "999")
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "999")
    paragraph._p.append(start)
    paragraph._p.append(end)


def _html_texts(html: str):
    """Текстовые узлы <p>/<h1>/<td> по порядку (для проверки порядка блоков)."""
    return re.findall(r"<(?:p|h1|h2|h3|td[^>]*)>(.*?)</(?:p|h1|h2|h3|td)>", html, re.S)


class WordPreviewRegression(unittest.TestCase):
    def _tmp_doc(self, tmp: Path, name: str = "t.docx") -> Path:
        return tmp / name

    def test_01_order_paragraph_table_paragraph(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            doc.add_paragraph("FIRST-BLOCK")
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "MID-TABLE"
            doc.add_paragraph("LAST-BLOCK")
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            texts = " ".join(_html_texts(html))
            self.assertLess(texts.index("FIRST-BLOCK"), texts.index("MID-TABLE"))
            self.assertLess(texts.index("MID-TABLE"), texts.index("LAST-BLOCK"))
            self.assertEqual(stats["tables"], 1)

    def test_02_external_hyperlink_kept(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            p = doc.add_paragraph()
            _add_hyperlink(p, "https://example.com/path?q=1", "LINK-TEXT")
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            self.assertIn('<a href="https://example.com/path?q=1">LINK-TEXT</a>', html)
            self.assertEqual(stats.get("links"), 1)

    def test_03_blocked_scheme_keeps_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            p = doc.add_paragraph()
            _add_hyperlink(p, "javascript:alert(1)", "EVIL-TEXT")
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            self.assertNotIn("javascript:", html)
            self.assertIn("EVIL-TEXT", html)
            self.assertIn("javascript", (stats.get("blockedUrls") or []))

    def test_04_internal_anchor_link(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            target = doc.add_paragraph("TARGET-PARA")
            _add_bookmark(target, "sec1")
            p = doc.add_paragraph()
            _add_anchor_link(p, "sec1", "GOTO-TEXT")
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            self.assertIn('id="bm-sec1"', html)
            self.assertIn('<a href="#bm-sec1">GOTO-TEXT</a>', html)

    def test_05_multilevel_numbering_values(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 100, [(1, "decimal", "%1."), (1, "decimal", "%1.%2."),
                                     (1, "lowerLetter", "%1.%2.%3.")], 100)
            for ilvl in (0, 1, 1, 2, 0):
                p = doc.add_paragraph(f"ITEM-{ilvl}")
                _set_num_pr(p, 100, ilvl)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["1.", "1.1.", "1.2.", "1.2.a.", "2."])

    def test_06_restart_new_num_starts_over(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 101, [(1, "decimal", "%1.")], 101)
            _add_numbering(doc, 101, [(1, "decimal", "%1.")], 102)
            for _ in range(3):
                p = doc.add_paragraph("A")
                _set_num_pr(p, 101, 0)
            for _ in range(2):
                p = doc.add_paragraph("B")
                _set_num_pr(p, 102, 0)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["1.", "2.", "3.", "1.", "2."])

    def test_07_numbering_via_style(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 102, [(5, "decimal", "%1.")], 103)
            style = doc.styles.add_style("SynNumStyle", WD_STYLE_TYPE.PARAGRAPH)
            pPr = style.element.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            nid = OxmlElement("w:numId")
            nid.set(qn("w:val"), "103")
            numPr.append(nid)
            pPr.append(numPr)
            doc.add_paragraph("STYLED-ONE", style="SynNumStyle")
            doc.add_paragraph("STYLED-TWO", style="SynNumStyle")
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["5.", "6."])

    def test_08_header_rendered(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            header = doc.sections[0].header
            hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            hp.text = "HEADER-TEXT"
            doc.add_paragraph("BODY-TEXT")
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            self.assertEqual(stats.get("headers"), 1)
            self.assertIn("HEADER-TEXT", html)
            self.assertIn('class="word-header"', html)

    def test_09_merged_cells_no_duplication_and_image_in_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "H-MERGED"
            table.cell(0, 0).merge(table.cell(0, 1))
            table.cell(0, 0).text = "H-MERGED"
            table.cell(1, 0).text = "V-TOP"
            # Вертикальное объединение корректной разметкой Word
            # (вручную raw-XML: первая ячейка restart с текстом,
            # вторая — продолжение vMerge без val и без текста).
            from docx.table import Table as _Table
            vtbl_el = parse_xml(
                f'<w:tbl {nsdecls("w")}><w:tblPr><w:tblW w:w="0" w:type="auto"/>'
                f"</w:tblPr><w:tblGrid><w:gridCol/></w:tblGrid>"
                f'<w:tr><w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr>'
                f"<w:p><w:r><w:t>V-MERGED</w:t></w:r></w:p></w:tc></w:tr>"
                f'<w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr>'
                f"<w:p><w:r><w:t></w:t></w:r></w:p></w:tc></w:tr></w:tbl>")
            doc.element.body.append(vtbl_el)
            _Table(vtbl_el, doc)  # проверка обёртки не падает
            pic_cell = table.cell(1, 0)
            pic_cell.paragraphs[0].add_run().add_picture(io.BytesIO(PIXEL_PNG))
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            self.assertEqual(html.count("H-MERGED"), 1)
            self.assertEqual(html.count("V-MERGED"), 1)
            self.assertIn('rowspan="2"', html)
            self.assertIn("<img", html)

    def test_10_cache_rehit_and_version(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            doc = Document()
            doc.add_paragraph("CACHE-ME")
            src = tmp / "c.docx"
            doc.save(str(src))
            cache = tmp / "cache"
            first = ensure_docx_preview(src, cache)
            self.assertFalse(first["cacheHit"])
            self.assertEqual(first.get("renderVersion"), DOCX_HTML_RENDER_VERSION)
            second = ensure_docx_preview(src, cache)
            self.assertTrue(second["cacheHit"])
            self.assertEqual(second["bytes"], first["bytes"])
            self.assertFalse((cache / "preview.lock").exists())

    def test_11_old_cache_invalidated(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            doc = Document()
            doc.add_paragraph("FRESH-CONTENT")
            src = tmp / "o.docx"
            doc.save(str(src))
            cache = tmp / "cache"
            cache.mkdir(parents=True, exist_ok=True)
            (cache / "preview.html").write_text("OLD-CONTENT", encoding="utf-8")
            (cache / "preview.json").write_text(json.dumps({
                "file": "preview.html", "bytes": 11,
                "sourceMtimeNs": src.stat().st_mtime_ns,
                "sourceSize": src.stat().st_size,
                "cacheHit": True}), encoding="utf-8")
            meta = ensure_docx_preview(src, cache)
            self.assertFalse(meta["cacheHit"])
            self.assertEqual(meta.get("renderVersion"), DOCX_HTML_RENDER_VERSION)
            self.assertNotIn("OLD-CONTENT", (cache / "preview.html").read_text(encoding="utf-8"))
            leftovers = list(cache.glob("preview.html.*.tmp"))
            self.assertEqual(leftovers, [])

    def test_12_concurrent_same_doc(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            doc = Document()
            doc.add_paragraph("CONC-PARA")
            src = tmp / "k.docx"
            doc.save(str(src))
            cache = tmp / "cache"

            def once(_):
                return ensure_docx_preview(src, cache)

            with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
                results = list(pool.map(once, range(8)))
            byted = {r["bytes"] for r in results}
            self.assertEqual(len(byted), 1)
            self.assertFalse((cache / "preview.lock").exists())


    def test_13_line_break_vs_literal(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            p = doc.add_paragraph()
            p.add_run("LINE-A")
            _add_break(p.add_run())
            p.add_run("LINE-B")
            doc.add_paragraph("TEXT-WITH-LITERAL <br> INSIDE")
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            self.assertIn("LINE-A<br>LINE-B", html)
            self.assertIn("TEXT-WITH-LITERAL &lt;br&gt; INSIDE", html)

    def test_14_break_inside_hyperlink(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            p = doc.add_paragraph()
            r_id = p.part.relate_to("https://example.com/x", RT.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), r_id)
            run = OxmlElement("w:r")
            txt = OxmlElement("w:t")
            txt.text = "CLICK-"
            run.append(txt)
            link.append(run)
            run2 = OxmlElement("w:r")
            txt2 = OxmlElement("w:t")
            txt2.text = "HERE"
            run2.append(txt2)
            run2.append(OxmlElement("w:br"))
            link.append(run2)
            p._p.append(link)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            self.assertIn('<a href="https://example.com/x">CLICK-HERE<br></a>', html)

    def test_15_restart_zero_never_restarts(self):
        # Случай A из рецензии: lvlRestart=0 — уровень 1 никогда
        # не перезапускается: 0/1/1/0/1 -> 1., 1.1., 1.2., 2., 2.3.
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 120, [(1, "decimal", "%1."), (1, "decimal", "%1.%2.")],
                           120, restarts={1: "0"})
            seq = [0, 1, 1, 0, 1]
            for ilvl in seq:
                p = doc.add_paragraph(f"R-{ilvl}")
                _set_num_pr(p, 120, ilvl)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["1.", "1.1.", "1.2.", "2.", "2.3."])

    def test_16_explicit_restart_on_level(self):
        # Случай C из рецензии: уровень 2, lvlRestart=1 (триггер ilvl 0):
        # уровень 0 вызывает перезапуск, уровень 1 — нет.
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 121, [(1, "decimal", "%1."), (1, "decimal", "%1.%2."),
                                     (1, "decimal", "%1.%2.%3.")],
                           121, restarts={2: "1"})
            seq = [0, 2, 1, 2]
            for ilvl in seq:
                p = doc.add_paragraph(f"E-{ilvl}")
                _set_num_pr(p, 121, ilvl)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["1.", "1.1.1.", "1.1.", "1.1.2."])

    def test_24_absent_restart_from_previous_level(self):
        # Случай D из рецензии: уровень 2 без lvlRestart перезапускается
        # при использовании предыдущего уровня (ilvl 1).
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 122, [(1, "decimal", "%1."), (1, "decimal", "%1.%2."),
                                     (1, "decimal", "%1.%2.%3.")], 122)
            seq = [0, 1, 2, 1, 2]
            for ilvl in seq:
                p = doc.add_paragraph(f"D-{ilvl}")
                _set_num_pr(p, 122, ilvl)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["1.", "1.1.", "1.1.1.", "1.2.", "1.2.1."])

    def test_25_restart_one_triggers_on_level_zero(self):
        # Случай B из рецензии: уровень 1, lvlRestart=1 (триггер ilvl 0):
        # после нового уровня 0 начинается заново.
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 123, [(1, "decimal", "%1."), (1, "decimal", "%1.%2.")],
                           123, restarts={1: "1"})
            seq = [0, 1, 1, 0, 1]
            for ilvl in seq:
                p = doc.add_paragraph(f"T-{ilvl}")
                _set_num_pr(p, 123, ilvl)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["1.", "1.1.", "1.2.", "2.", "2.1."])

    def test_26_level_zero_never_self_restarts(self):
        # Случай E из рецензии: последовательные пункты уровня 0
        # не превращаются в повторяющиеся «1.» (нет сброса текущего уровня).
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 124, [(1, "decimal", "%1.")], 124,
                           restarts={0: "1"})
            for _ in range(3):
                p = doc.add_paragraph("Z")
                _set_num_pr(p, 124, 0)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["1.", "2.", "3."])

    def test_17_override_level_redefines_format(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 112, [(1, "decimal", "%1.")], 112,
                           overrides={0: 1},
                           override_levels={0: ("lowerLetter", "(%1)")})
            for _ in range(2):
                doc.add_paragraph("OV")
                _set_num_pr(doc.paragraphs[-1], 112, 0)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["(a)", "(b)"])

    def test_18_partial_direct_numpr_with_style(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 113, [(1, "decimal", "%1."), (1, "decimal", "%1.%2.")], 113)
            style = doc.styles.add_style("SynPartial", WD_STYLE_TYPE.PARAGRAPH)
            pPr = style.element.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            nid = OxmlElement("w:numId")
            nid.set(qn("w:val"), "113")
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), "1")
            numPr.append(il)
            numPr.append(nid)
            pPr.append(numPr)
            p = doc.add_paragraph("PARTIAL-ONE", style="SynPartial")
            _set_num_id_only(p, 113)  # прямой numId без ilvl + ilvl из стиля
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["1.1."])

    def test_19_header_uses_own_relationships(self):
        with tempfile.TemporaryDirectory() as tmp:
            body_png = _png_variant(b"body-image")
            head_png = _png_variant(b"header-image")
            doc = Document()
            bp = doc.add_paragraph()
            _add_hyperlink(bp, "https://body.example/x", "BODY-LINK")
            bp.add_run().add_picture(io.BytesIO(body_png))
            header = doc.sections[0].header
            hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            _add_hyperlink(hp, "https://header.example/y", "HEADER-LINK")
            hp.add_run().add_picture(io.BytesIO(head_png))
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            self.assertIn('<a href="https://body.example/x">BODY-LINK</a>', html)
            self.assertIn('<a href="https://header.example/y">HEADER-LINK</a>', html)
            body_b64 = base64.b64encode(body_png).decode()
            head_b64 = base64.b64encode(head_png).decode()
            self.assertNotEqual(body_b64, head_b64)
            cut = html.find('class="word-header"')
            self.assertGreater(cut, 0)
            body_part, head_part = html[:cut], html[cut:]
            self.assertIn(body_b64, body_part)
            self.assertNotIn(head_b64, body_part)
            self.assertIn(head_b64, head_part)
            self.assertNotIn(body_b64, head_part)

    def test_20_first_page_header_flagged(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            header = doc.sections[0].header
            hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            hp.text = "DEFAULT-HDR"
            doc.sections[0].different_first_page_header_footer = True
            sectPr = doc.sections[0]._sectPr
            default_ref = sectPr.find(qn("w:headerReference"))
            first_ref = OxmlElement("w:headerReference")
            first_ref.set(qn("w:type"), "first")
            first_ref.set(qn("r:id"), default_ref.get(qn("r:id")))
            sectPr.append(first_ref)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            self.assertIn("DEFAULT-HDR", html)
            self.assertIn("first", stats.get("headersUnsupported", []))

    def test_21_stale_cache_concurrent_regeneration(self):
        import threading
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            doc = Document()
            doc.add_paragraph("STALE-PARA")
            src = tmp / "s.docx"
            doc.save(str(src))
            cache = tmp / "cache"
            cache.mkdir(parents=True, exist_ok=True)
            (cache / "preview.html").write_text("STALE-HTML", encoding="utf-8")
            (cache / "preview.json").write_text(json.dumps({
                "file": "preview.html", "bytes": 10,
                "sourceMtimeNs": src.stat().st_mtime_ns,
                "sourceSize": src.stat().st_size,
                "cacheHit": True}), encoding="utf-8")
            entered = threading.Event()
            release = threading.Event()
            locked = threading.Event()
            waiting = threading.Event()
            calls = {"n": 0}
            real = engine_word.docx_to_html_string
            real_read = engine_word._read_valid_meta
            real_acquire = engine_word._try_acquire_lock

            def slow(src_path):
                entered.set()
                self.assertTrue(release.wait(timeout=30))
                return real(src_path)

            def counting_read(meta_path, html_path, stat):
                calls["n"] += 1
                if locked.is_set() and calls["n"] >= 4:
                    waiting.set()
                return real_read(meta_path, html_path, stat)

            def acquiring(cache_dir):
                ok = real_acquire(cache_dir)
                if ok:
                    locked.set()
                return ok

            errors = []
            import unittest.mock as mock
            with mock.patch.object(engine_word, "docx_to_html_string", side_effect=slow), \
                 mock.patch.object(engine_word, "_read_valid_meta", side_effect=counting_read), \
                 mock.patch.object(engine_word, "_try_acquire_lock", side_effect=acquiring):
                def first():
                    try:
                        return ("ok", ensure_docx_preview(src, cache))
                    except Exception as exc:  # noqa: BLE001
                        errors.append(exc)
                        return ("err", exc)

                with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
                    fut1 = pool.submit(first)
                    self.assertTrue(entered.wait(timeout=30))
                    self.assertTrue(locked.wait(timeout=30))
                    fut2 = pool.submit(first)
                    # Второй поток подтверждённо вошёл в ожидание чужой
                    # генерации — только затем отпускаем первый.
                    self.assertTrue(waiting.wait(timeout=30))
                    release.set()
                    kind1, res1 = fut1.result(timeout=60)
                    kind2, res2 = fut2.result(timeout=60)
            self.assertEqual(errors, [])
            self.assertEqual((kind1, kind2), ("ok", "ok"))
            self.assertEqual(res1["bytes"], res2["bytes"])
            for res in (res1, res2):
                self.assertEqual(res.get("renderVersion"), DOCX_HTML_RENDER_VERSION)
            self.assertNotIn("STALE-HTML", (cache / "preview.html").read_text(encoding="utf-8"))

    def test_22_corrupt_cache_regenerates(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            doc = Document()
            doc.add_paragraph("RECOVER-ME")
            src = tmp / "r.docx"
            doc.save(str(src))
            cache = tmp / "cache"
            cache.mkdir(parents=True, exist_ok=True)
            (cache / "preview.html").write_text("", encoding="utf-8")
            (cache / "preview.json").write_text("{not-json", encoding="utf-8")
            meta = ensure_docx_preview(src, cache)
            self.assertFalse(meta["cacheHit"])
            self.assertGreater((cache / "preview.html").stat().st_size, 0)
            again = ensure_docx_preview(src, cache)
            self.assertTrue(again["cacheHit"])

    def test_23_explicit_formatting_off(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            p = doc.add_paragraph()
            r1 = p.add_run("BOLD-OFF")
            _set_run_off(r1, "b", "false")
            r2 = p.add_run("ITALIC-OFF")
            _set_run_off(r2, "i", "off")
            r3 = p.add_run("ULINE-NONE")
            _set_run_off(r3, "u", "none")
            r4 = p.add_run("BOLD-ON")
            rpr = r4._r.get_or_add_rPr()
            rpr.append(OxmlElement("w:b"))
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, _ = docx_to_html_string(src)
            self.assertNotIn("<b>BOLD-OFF</b>", html)
            self.assertNotIn("<i>ITALIC-OFF</i>", html)
            self.assertIn("BOLD-OFF", html)
            self.assertIn("ITALIC-OFF", html)
            self.assertIn("ULINE-NONE", html)
            self.assertNotIn("<u>ULINE-NONE</u>", html)
            self.assertIn("<b>BOLD-ON</b>", html)


    def test_27_invalid_trigger_never_restarts(self):
        # Недопустимый для уровня триггер (уровень 1, lvlRestart=3 —
        # триггер ilvl 2 глубже текущего): сброса нет, запись в диагностике.
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            _add_numbering(doc, 125, [(1, "decimal", "%1."), (1, "decimal", "%1.%2."),
                                     (1, "decimal", "%1.%2.%3.")],
                           125, restarts={1: "3"})
            seq = [0, 1, 2, 1]
            for ilvl in seq:
                p = doc.add_paragraph(f"X-{ilvl}")
                _set_num_pr(p, 125, ilvl)
            src = self._tmp_doc(Path(tmp))
            doc.save(str(src))
            html, stats = docx_to_html_string(src)
            marks = re.findall(r'<span class="wnum">(.*?)</span>', html)
            self.assertEqual(marks, ["1.", "1.1.", "1.1.1.", "1.2."])
            diags = stats.get("numberingUnsupported", [])
            self.assertTrue(any("lvlRestart:3" in d for d in diags), diags)


if __name__ == "__main__":
    unittest.main()
