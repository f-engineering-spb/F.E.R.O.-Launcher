"""S02 Excel fast preview module (stand-alone, launcher-connectable).

Provides the fastest readable Excel preview routes WITHOUT touching the
working launcher:

- ``get_fast_html``: first-viewport lite HTML (plain values, no styles).
- ``get_fast_png``:  direct PIL thumbnail PNG from the same values (no COM,
  no PDF, no browser snapshot).

Both use a permanent file cache keyed by ``(source mtime_ns, size, purpose,
viewport)`` so:

- repeat open is served from cache (no rebuild),
- repeat after process restart is served from cache (nothing in-memory),
- source change yields a new key and rebuilds.

Only ``openpyxl`` (already required) + ``Pillow`` are needed for the PNG
route.  No formula evaluator is implemented: ``data_only=True`` cached
values are used as-is; formulas without a stored value render empty
(same documented loss as the existing HTML route).

Designed so the launcher can later call ``get_fast_html`` /
``get_fast_png`` from ``server.py`` and serve the files under
``/cache/excel/...`` with the existing epoch/late-response guards.
"""

from __future__ import annotations

import hashlib
import html
import io
import json
import time
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # pragma: no cover
    Image = None

CACHE_VERSION = "s02-v1"
DEFAULT_MAX_ROWS = 50
DEFAULT_MAX_COLS = 20

_PNG_FONT_CANDIDATES = (
    "C:/Windows/Fonts/arial.ttf",
    "C:/Windows/Fonts/segoeui.ttf",
)


def _font(size: int):
    for candidate in _PNG_FONT_CANDIDATES:
        try:
            return ImageFont.truetype(candidate, size)
        except Exception:
            continue
    return ImageFont.load_default()


def cache_key_for(path: Path, purpose: str, extra: str = "") -> str:
    """Stable key for one exact workbook revision + viewport."""
    stat = path.stat()
    raw = "|".join(
        (
            CACHE_VERSION,
            purpose,
            str(path.resolve()),
            str(stat.st_mtime_ns),
            str(stat.st_size),
            extra,
        )
    )
    return hashlib.sha1(raw.casefold().encode("utf-8")).hexdigest()[:20]


def cache_dir_for(path: Path, purpose: str, base_dir: Path, extra: str = "") -> Path:
    return base_dir / cache_key_for(path, purpose, extra)


def _sheet_names_fast(path: Path) -> list[str] | None:
    """Sheet names via ZIP central read (no workbook load)."""
    import zipfile
    import xml.etree.ElementTree as ET

    try:
        if not zipfile.is_zipfile(path):
            return None
        with zipfile.ZipFile(path, "r") as zf:
            if "xl/workbook.xml" not in zf.namelist():
                return None
            root = ET.fromstring(zf.read("xl/workbook.xml"))
            ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            names = [el.attrib.get("name", "") for el in root.findall(".//main:sheet", ns)]
            if names and all(names):
                return names
    except Exception:
        return None
    return None


def read_sheet_names(path: Path) -> list[str]:
    """Return sheet names via fast zip inspection or openpyxl fallback."""
    src = Path(path)
    names = _sheet_names_fast(src)
    if names is not None:
        return names
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel fast preview")
    wb_probe = openpyxl.load_workbook(
        filename=str(src), read_only=True, data_only=True
    )
    try:
        return list(wb_probe.sheetnames)
    finally:
        wb_probe.close()


def _serialize_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.strftime("%d.%m.%Y")
        return value.strftime("%d.%m.%Y %H:%M")
    text = str(value)
    return text.strip()


def read_values_fast(
    path: Path,
    sheet_index: int = 0,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_cols: int = DEFAULT_MAX_COLS,
) -> dict:
    """Read first-viewport values with streaming ``read_only=True``.

    No styles, no merges, no formula evaluation.  Raises on missing file
    or bad sheet index.
    """
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel fast preview")
    src = Path(path)
    if not src.is_file():
        raise FileNotFoundError(f"Excel file not found: {src}")

    names = _sheet_names_fast(src)
    if names is None:
        wb_probe = openpyxl.load_workbook(
            filename=str(src), read_only=True, data_only=True
        )
        try:
            names = list(wb_probe.sheetnames)
        finally:
            wb_probe.close()
    if sheet_index < 0 or sheet_index >= len(names):
        raise IndexError(f"Sheet {sheet_index} out of range ({len(names)} sheets)")

    wb = openpyxl.load_workbook(filename=str(src), read_only=True, data_only=True)
    try:
        ws = wb[names[sheet_index]]
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = [_serialize_value(v) for v in list(row)[:max_cols]]
            while cells and cells[-1] == "":
                cells.pop()
            if any(cells):
                rows.append(cells)
            if len(rows) >= max_rows:
                break
        truncated = len(rows) >= max_rows
    finally:
        wb.close()

    width = 0
    for r in rows:
        width = max(width, len(r))
    grid = [r + [""] * (width - len(r)) for r in rows]
    return {
        "sheet": names[sheet_index],
        "sheets": names,
        "grid": grid,
        "rows": len(grid),
        "columns": width,
        "truncated": truncated,
    }


def render_fast_html(values: dict, source_name: str) -> str:
    title = values.get("sheet", source_name)
    body_rows = []
    for i, row in enumerate(values.get("grid", [])):
        tag = "th" if i == 0 else "td"
        cells = "".join(f"<{tag}>{html.escape(c)}</{tag}>" for c in row)
        body_rows.append(f"<tr>{cells}</tr>")
    notice = (
        f"Первые {DEFAULT_MAX_ROWS} строк × {DEFAULT_MAX_COLS} столбцов. "
        "Полный файл — в Excel."
        if values.get("truncated")
        else "Быстрый просмотр первого вида"
    )
    return (
        "<!doctype html><html lang='ru'><head><meta charset='utf-8'>"
        f"<title>{html.escape(str(title))}</title>"
        "<style>body{margin:0;font:12px 'Segoe UI',Arial,sans-serif;color:#20262d}"
        ".n{position:sticky;top:0;background:#f2f6f9;border-bottom:1px solid #d3dde4;"
        "padding:6px 10px;color:#607080}table{border-collapse:collapse}"
        "td,th{border:1px solid #cbd5dc;padding:3px 8px;white-space:nowrap;max-width:320px;"
        "overflow:hidden;text-overflow:ellipsis}"
        "tr:first-child th{background:#eef3f6}</style></head><body>"
        f"<div class=n>{html.escape(notice)}</div>"
        f"<table>{''.join(body_rows)}</table></body></html>"
    )


def render_fast_png(values: dict, source_name: str) -> bytes:
    """Direct thumbnail PNG from values (no COM/PDF/browser)."""
    if Image is None:
        raise RuntimeError("Pillow is required for Excel fast PNG")
    grid = values.get("grid", []) or [["(пусто)"]]
    n_cols = max(len(r) for r in grid)
    grid = [r + [""] * (n_cols - len(r)) for r in grid]

    font = _font(15)
    header_font = _font(16)
    pad_x, row_h, header_h = 10, 26, 30
    widths = []
    drawer_probe = ImageDraw.Draw(Image.new("RGB", (8, 8)))
    for c in range(n_cols):
        best = 60
        for r in grid:
            box = drawer_probe.textbbox((0, 0), r[c] or " ", font=font)
            best = max(best, box[2] - box[0] + pad_x * 2)
        widths.append(min(best, 340))
    table_w = sum(widths) + 1
    table_h = header_h + row_h * max(0, len(grid) - 1) + 1
    title_h = 34
    img = Image.new("RGB", (table_w + 20, table_h + title_h + 20), "white")
    draw = ImageDraw.Draw(img)
    draw.text(
        (10, 8),
        f"{source_name} — {values.get('sheet', '')}",
        fill="#34424e",
        font=header_font,
    )
    y0 = title_h + 10
    x0 = 10
    for r_idx, row in enumerate(grid):
        h = header_h if r_idx == 0 else row_h
        x = x0
        for c_idx, cell in enumerate(row):
            w = widths[c_idx]
            fill = "#eef3f6" if r_idx == 0 else "white"
            draw.rectangle([x, y0, x + w, y0 + h], fill=fill, outline="#cbd5dc")
            draw.text(
                (x + pad_x, y0 + 5),
                (cell[:60] + "…") if len(cell) > 61 else cell,
                fill="#20262d",
                font=header_font if r_idx == 0 else font,
            )
            x += w
        y0 += h
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _write_manifest(out_dir: Path, payload: dict) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "manifest.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def _cache_url(out: Path, base_dir: Path, key: str, filename: str) -> str:
    try:
        parts = list(out.resolve().parts)
        if "cache" in parts:
            idx = parts.index("cache")
            return "/" + "/".join(parts[idx:])
    except Exception:
        pass
    return f"/cache/excel/{key}/{filename}"


def get_fast_html(
    path: Path,
    base_dir: Path,
    sheet_index: int = 0,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_cols: int = DEFAULT_MAX_COLS,
) -> dict:
    """Cached lite-HTML preview. Returns timings + cache flag."""
    src = Path(path)
    extra = f"s{sheet_index}_{max_rows}x{max_cols}"
    key = cache_key_for(src, "fast-html", extra)
    out_dir = Path(base_dir) / key
    out = out_dir / "preview.html"
    meta = out_dir / "manifest.json"
    url = _cache_url(out, base_dir, key, "preview.html")
    if out.exists() and out.stat().st_size > 0 and meta.exists():
        try:
            saved = json.loads(meta.read_text(encoding="utf-8"))
            if (
                saved.get("sourcePath") == str(src.resolve())
                and saved.get("sourceMtimeNs") == src.stat().st_mtime_ns
                and saved.get("sourceSize") == src.stat().st_size
                and saved.get("cacheKey") == key
            ):
                t0 = time.perf_counter()
                blob = out.read_bytes()
                serve_ms = (time.perf_counter() - t0) * 1000
                return {
                    "path": str(out),
                    "url": url,
                    "bytes": len(blob),
                    "cacheHit": True,
                    "cacheKey": key,
                    "engineMs": 0.0,
                    "serveMs": serve_ms,
                    "sheetIndex": sheet_index,
                    **saved.get("info", {}),
                }
        except (OSError, ValueError):
            pass
    t0 = time.perf_counter()
    values = read_values_fast(src, sheet_index, max_rows, max_cols)
    page = render_fast_html(values, src.name)
    engine_ms = (time.perf_counter() - t0) * 1000
    out_dir.mkdir(parents=True, exist_ok=True)
    out.write_text(page, encoding="utf-8")
    info = {
        "sheet": values["sheet"],
        "sheetIndex": sheet_index,
        "rows": values["rows"],
        "columns": values["columns"],
        "truncated": values["truncated"],
    }
    _write_manifest(
        out_dir,
        {
            "sourcePath": str(src.resolve()),
            "sourceMtimeNs": src.stat().st_mtime_ns,
            "sourceSize": src.stat().st_size,
            "cacheKey": key,
            "kind": "fast-html",
            "viewport": extra,
            "info": info,
        },
    )
    t1 = time.perf_counter()
    blob = out.read_bytes()
    serve_ms = (time.perf_counter() - t1) * 1000
    return {
        "path": str(out),
        "url": url,
        "bytes": len(blob),
        "cacheHit": False,
        "cacheKey": key,
        "engineMs": engine_ms,
        "serveMs": serve_ms,
        **info,
    }


def get_fast_workbook(
    path: Path,
    base_dir: Path,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_cols: int = DEFAULT_MAX_COLS,
) -> dict:
    """Quickly inspect sheet names and prepare the first sheet preview."""
    src = Path(path)
    if not src.is_file():
        raise FileNotFoundError(f"Excel file not found: {src}")
    names = read_sheet_names(src)
    sheets = [{"index": idx, "name": name} for idx, name in enumerate(names)]
    sheet0 = get_fast_html(src, base_dir, sheet_index=0, max_rows=max_rows, max_cols=max_cols)
    return {
        "name": src.name,
        "path": str(src.resolve()),
        "sourcePath": str(src.resolve()),
        "sourceName": src.name,
        "sourceType": src.suffix.lstrip(".").upper(),
        "sheets": sheets,
        "activeSheetIndex": 0,
        "thumbnailUrl": sheet0["url"],
        "sheetUrl": sheet0["url"],
        "engineMs": sheet0["engineMs"],
        "serveMs": sheet0["serveMs"],
        "cacheHit": sheet0["cacheHit"],
        "cacheKey": sheet0.get("cacheKey", ""),
    }


def get_fast_png(
    path: Path,
    base_dir: Path,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_cols: int = DEFAULT_MAX_COLS,
) -> dict:
    """Cached direct-PNG preview. Returns timings + cache flag."""
    src = Path(path)
    extra = f"{max_rows}x{max_cols}"
    key = cache_key_for(src, "fast-png", extra)
    out_dir = Path(base_dir) / key
    out = out_dir / "preview.png"
    meta = out_dir / "manifest.json"
    if out.exists() and out.stat().st_size > 0 and meta.exists():
        try:
            saved = json.loads(meta.read_text(encoding="utf-8"))
            if (
                saved.get("sourcePath") == str(src.resolve())
                and saved.get("sourceMtimeNs") == src.stat().st_mtime_ns
                and saved.get("sourceSize") == src.stat().st_size
                and saved.get("cacheKey") == key
            ):
                t0 = time.perf_counter()
                blob = out.read_bytes()
                serve_ms = (time.perf_counter() - t0) * 1000
                return {
                    "path": str(out),
                    "url": f"/cache/excel/{key}/preview.png",
                    "bytes": len(blob),
                    "cacheHit": True,
                    "engineMs": 0.0,
                    "serveMs": serve_ms,
                    **saved.get("info", {}),
                }
        except (OSError, ValueError):
            pass
    t0 = time.perf_counter()
    values = read_values_fast(src, 0, max_rows, max_cols)
    blob = render_fast_png(values, src.name)
    engine_ms = (time.perf_counter() - t0) * 1000
    out_dir.mkdir(parents=True, exist_ok=True)
    out.write_bytes(blob)
    info = {
        "sheet": values["sheet"],
        "rows": values["rows"],
        "columns": values["columns"],
        "truncated": values["truncated"],
    }
    _write_manifest(
        out_dir,
        {
            "sourcePath": str(src.resolve()),
            "sourceMtimeNs": src.stat().st_mtime_ns,
            "sourceSize": src.stat().st_size,
            "cacheKey": key,
            "kind": "fast-png",
            "viewport": extra,
            "info": info,
        },
    )
    t1 = time.perf_counter()
    blob = out.read_bytes()
    serve_ms = (time.perf_counter() - t1) * 1000
    return {
        "path": str(out),
        "url": f"/cache/excel/{key}/preview.png",
        "bytes": len(blob),
        "cacheHit": False,
        "engineMs": engine_ms,
        "serveMs": serve_ms,
        **info,
    }
