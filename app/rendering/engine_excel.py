"""Движок Excel: один разбор книги → все вкладки мгновенно.

Ключевое исправление против старого кода server.py::excel_sheet_html:
раньше КАЖДЫЙ клик по вкладке делал полный load_workbook(read_only=False)
всей книги (~6 с на смете 833 КБ). Теперь книга разбирается ОДИН раз,
все листы рендерятся пакетно, переключение вкладок — дисковый кэш-хит.

Совместимость: имена файлов sheet-{N}-v10.html/.json, ключи метаданных
и HTML-оболочка iframe (app.js: zoom/rotate/fit/hand) сохранены 1:1.
"""

from __future__ import annotations

import html as _html
import json
import time
from collections import OrderedDict
from pathlib import Path

MAX_ROWS = 2000
MAX_COLS = 100
SHEET_VERSION = "v10"

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlrd
except ImportError:
    xlrd = None

from ._sheet_shell import render_sheet_page

# Процессный LRU разобранных книг: повторные вкладки без перепарсинга.
_BOOK_CACHE: OrderedDict = OrderedDict()
_BOOK_CACHE_SIZE = 3


def cell_color(color: object) -> str | None:
    value = getattr(color, "rgb", None)
    value = str(value) if value is not None else ""
    if not value or len(value) < 6 or (len(value) == 8 and value[:2] == "00"):
        return None
    return f"#{value[-6:]}"


def _book_cache_key(path: Path) -> tuple:
    stat = path.stat()
    return (str(path.resolve()), stat.st_mtime_ns, stat.st_size)


def _load_book_cached(xlsx_path: Path):
    """Вернуть открытую книгу, переиспользовав LRU при неизменном файле."""
    key = _book_cache_key(xlsx_path)
    book = _BOOK_CACHE.get(key)
    if book is not None:
        _BOOK_CACHE.move_to_end(key)
        return book, True
    mine = str(xlsx_path.resolve())
    for old_key in [k for k in _BOOK_CACHE if k[0] == mine]:
        try:
            _BOOK_CACHE[old_key].close()
        except Exception:
            pass
        del _BOOK_CACHE[old_key]
    while len(_BOOK_CACHE) >= _BOOK_CACHE_SIZE:
        _, old_book = _BOOK_CACHE.popitem(last=False)
        try:
            old_book.close()
        except Exception:
            pass
    book = openpyxl.load_workbook(xlsx_path, read_only=False, data_only=True)
    _BOOK_CACHE[key] = book
    return book, False


def get_sheet_names(xlsx_path: Path) -> list[str]:
    """Мгновенный список вкладок: прямой ZIP-парсинг xl/workbook.xml."""
    import xml.etree.ElementTree as ET
    import zipfile
    if zipfile.is_zipfile(xlsx_path):
        try:
            with zipfile.ZipFile(xlsx_path, "r") as z:
                tree = ET.fromstring(z.read("xl/workbook.xml"))
                ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                names = [el.attrib.get("name", "") for el in tree.findall(".//main:sheet", ns)]
                if names and all(names):
                    return names
        except Exception:
            pass
    book = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def _render_sheet_from_ws(sheet, sheet_index: int) -> tuple[str, dict]:
    """Построить (page, metadata) для одного листа уже открытой книги."""
    max_row = int(sheet.max_row or 1)
    max_column = int(sheet.max_column or 1)
    capped_col = min(max_column, MAX_COLS)
    capped_row = min(max_row, MAX_ROWS)
    columns = [
        c for c in range(1, capped_col + 1)
        if not sheet.column_dimensions[openpyxl.utils.get_column_letter(c)].hidden
    ]
    column_pixels = {}
    for c in columns:
        width = sheet.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 8.43
        column_pixels[c] = max(4, min(720, round(width * 7 + 5)))

    merged_start: dict = {}
    merged_skip: set = set()
    for area in sheet.merged_cells.ranges:
        shown = [c for c in columns if area.min_col <= c <= area.max_col]
        if not shown:
            continue
        start = (area.min_row, shown[0])
        merged_start[start] = (len(shown), area.max_row - area.min_row + 1)
        for r in range(area.min_row, area.max_row + 1):
            for c in shown:
                if (r, c) != start:
                    merged_skip.add((r, c))

    # ОДИН проход: читаем сетку ячеек единожды, дальше — скан и рендер из кэша.
    grid = list(sheet.iter_rows(min_row=1, max_row=capped_row, max_col=capped_col))
    col_index = {c: c - 1 for c in range(1, capped_col + 1)}

    first_value_row = 0
    last_value_row = 0
    for r in range(1, capped_row + 1):
        cells = grid[r - 1]
        hit = False
        for c in columns:
            if cells[col_index[c]].value is not None:
                hit = True
                break
        if hit:
            if not first_value_row:
                first_value_row = r
            last_value_row = r
    last_merge_row = max((a.max_row for a in sheet.merged_cells.ranges), default=0)
    rendered_rows = min(max(last_value_row, last_merge_row), MAX_ROWS)
    was_limited = max_row > MAX_ROWS

    first_rendered_row = first_value_row or 1
    rows = []
    for r in range(first_rendered_row, rendered_rows + 1):
        if sheet.row_dimensions[r].hidden:
            continue
        cells = grid[r - 1]
        out_cells = []
        for c in columns:
            if (r, c) in merged_skip:
                continue
            cell = cells[col_index[c]]
            raw = cell.value
            value = "" if raw is None else str(raw)
            css = []
            if cell.font.bold:
                css.append("font-weight:500")
            if cell.font.italic:
                css.append("font-style:italic")
            if cell.font.sz and round(cell.font.sz) != 11:
                css.append(f"font-size:{max(6, min(32, cell.font.sz))}px")
            if cell.font.name and cell.font.name.casefold() not in {"calibri", "arial", "segoe ui"}:
                css.append(f"font-family:{_html.escape(cell.font.name, quote=True)}")
            font_color = cell_color(cell.font.color)
            if font_color:
                css.append(f"color:{font_color}")
            fill = cell_color(cell.fill.fgColor)
            if cell.fill.fill_type == "solid" and fill:
                css.append(f"background:{fill}")
            if cell.alignment.horizontal in {"left", "center", "right"}:
                css.append(f"text-align:{cell.alignment.horizontal}")
            if cell.alignment.vertical in {"top", "center", "bottom"}:
                css.append(f"vertical-align:{cell.alignment.vertical}")
            if cell.alignment.wrap_text is False:
                css.append("white-space:pre;overflow:hidden")
            colspan, rowspan = merged_start.get((r, c), (1, 1))
            span = (f' colspan="{colspan}"' if colspan > 1 else "") + (
                f' rowspan="{rowspan}"' if rowspan > 1 else "")
            out_cells.append(f'<td{span} style="{";".join(css)}">{_html.escape(value)}</td>')
        authored_height = sheet.row_dimensions[r].height
        row_style = (f' style="height:{max(1, round(authored_height * 1.33))}px"'
                     if authored_height else "")
        rows.append(f"<tr{row_style}><th>{r}</th>{''.join(out_cells)}</tr>")

    cols = '<col class="row-number">' + "".join(
        f'<col style="width:{column_pixels[c]}px">' for c in columns)
    table_width = 38 + sum(column_pixels.values())
    notice = (f"Показаны первые {MAX_ROWS:,} строк. Полный рабочий файл откройте в Excel."
              if was_limited else "Просмотр без редактирования")
    metadata = {
        "name": sheet.title,
        "rows": max(0, rendered_rows - first_rendered_row + 1),
        "firstRow": first_rendered_row,
        "columns": len(columns),
        "limited": was_limited,
    }
    page = render_sheet_page(title=sheet.title, cols=cols, rows=rows,
                             table_width=table_width, notice=notice)
    return page, metadata


def ensure_workbook_rendered(xlsx_path: Path, cache_dir: Path) -> list[dict]:
    """Отрендерить ВСЕ листы за один разбор книги. Возвращает метаданные по индексам.

    Повторный вызов при неизменном файле — чистый дисковый кэш-хит
    (openpyxl не трогается вообще).
    """
    if openpyxl is None:
        raise RuntimeError("Для HTML-просмотра Excel нужен пакет openpyxl")
    cache_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = cache_dir / f"workbook-{SHEET_VERSION}.json"
    stat = xlsx_path.stat()
    fresh = False
    if manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            fresh = (
                manifest.get("sourcePath") == str(xlsx_path)
                and manifest.get("sourceMtimeNs") == stat.st_mtime_ns
                and manifest.get("sourceSize") == stat.st_size
            )
        except (OSError, json.JSONDecodeError):
            fresh = False

    book = None
    if fresh:
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            metas = []
            ok = True
            for i in range(manifest.get("sheets", 0)):
                mp = cache_dir / f"sheet-{i + 1}-{SHEET_VERSION}.json"
                hp = cache_dir / f"sheet-{i + 1}-{SHEET_VERSION}.html"
                if not (mp.exists() and hp.exists()):
                    ok = False
                    break
                metas.append(json.loads(mp.read_text(encoding="utf-8")))
            if ok:
                return metas
        except (OSError, json.JSONDecodeError):
            pass
    book, _ = _load_book_cached(xlsx_path)

    t0 = time.perf_counter()
    try:
        names = list(book.sheetnames)
        metas: list[dict] = []
        for i, name in enumerate(names):
            page, metadata = _render_sheet_from_ws(book[name], i)
            (cache_dir / f"sheet-{i + 1}-{SHEET_VERSION}.html").write_text(page, encoding="utf-8")
            (cache_dir / f"sheet-{i + 1}-{SHEET_VERSION}.json").write_text(
                json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
            metas.append(metadata)
        manifest_path.write_text(json.dumps({
            "sourcePath": str(xlsx_path),
            "sourceMtimeNs": stat.st_mtime_ns,
            "sourceSize": stat.st_size,
            "sheets": len(names),
            "renderMs": round((time.perf_counter() - t0) * 1000, 1),
        }, ensure_ascii=False), encoding="utf-8")
        return metas
    finally:
        pass


def render_sheet(xlsx_path: Path, sheet_index: int, cache_dir: Path) -> tuple[str, dict]:
    """Совместимый точечный рендер одного листа (кэш всех вкладок греется пакетно)."""
    metas = ensure_workbook_rendered(xlsx_path, cache_dir)
    if not (0 <= sheet_index < len(metas)):
        raise IndexError(f"Нет листа {sheet_index}: всего {len(metas)}")
    page = (cache_dir / f"sheet-{sheet_index + 1}-{SHEET_VERSION}.html").read_text(encoding="utf-8")
    return page, metas[sheet_index]


def ensure_xls_rendered(xls_path: Path, cache_dir: Path) -> list[dict]:
    """Быстрый путь для бинарных BIFF8 (.xls) через xlrd — ~10 мс, без COM."""
    if xlrd is None:
        raise RuntimeError("Для просмотра XLS нужен пакет xlrd")
    cache_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = cache_dir / f"workbook-{SHEET_VERSION}.json"
    stat = xls_path.stat()
    if manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            if (manifest.get("sourceMtimeNs") == stat.st_mtime_ns
                    and manifest.get("sourceSize") == stat.st_size
                    and manifest.get("engine") == "xlrd"):
                metas = []
                ok = True
                for i in range(manifest.get("sheets", 0)):
                    mp = cache_dir / f"sheet-{i + 1}-{SHEET_VERSION}.json"
                    if not mp.exists():
                        ok = False
                        break
                    metas.append(json.loads(mp.read_text(encoding="utf-8")))
                if ok:
                    return metas
        except (OSError, json.JSONDecodeError):
            pass

    t0 = time.perf_counter()
    book = xlrd.open_workbook(str(xls_path), on_demand=False)
    try:
        metas: list[dict] = []
        for i, name in enumerate(book.sheet_names()):
            sh = book.sheet_by_index(i)
            nrows = min(sh.nrows, MAX_ROWS)
            ncols = min(sh.ncols, MAX_COLS)
            rows = []
            for r in range(nrows):
                cells = []
                for c in range(ncols):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK,
                                      xlrd.XL_CELL_ERROR):
                        s = ""
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        s = "ИСТИНА" if v else "ЛОЖЬ"
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            s = xlrd.xldate_as_datetime(v, book.datemode).strftime("%d.%m.%Y")
                        except Exception:
                            s = str(v)
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        s = str(int(v)) if float(v).is_integer() else str(round(float(v), 4))
                    else:
                        s = str(v)
                    style = ' style="font-weight:500"' if r == 0 and s else ""
                    cells.append(f"<td{style}>{_html.escape(s)}</td>")
                rows.append(f"<tr><th>{r + 1}</th>{''.join(cells)}</tr>")
            cols = '<col class="row-number">' + "".join(
                '<col style="width:90px">' for _ in range(ncols))
            metadata = {"name": name, "rows": nrows, "firstRow": 1,
                        "columns": ncols, "limited": sh.nrows > MAX_ROWS}
            notice = ("Просмотр без редактирования"
                      if not metadata["limited"]
                      else f"Показаны первые {MAX_ROWS:,} строк. Полный файл откройте в Excel.")
            page = render_sheet_page(title=name, cols=cols, rows=rows,
                                     table_width=38 + 90 * ncols, notice=notice)
            (cache_dir / f"sheet-{i + 1}-{SHEET_VERSION}.html").write_text(page, encoding="utf-8")
            (cache_dir / f"sheet-{i + 1}-{SHEET_VERSION}.json").write_text(
                json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
            metas.append(metadata)
        manifest_path.write_text(json.dumps({
            "engine": "xlrd",
            "sourcePath": str(xls_path),
            "sourceMtimeNs": stat.st_mtime_ns,
            "sourceSize": stat.st_size,
            "sheets": len(metas),
            "renderMs": round((time.perf_counter() - t0) * 1000, 1),
        }, ensure_ascii=False), encoding="utf-8")
        return metas
    finally:
        book.release_resources()
